# ch19_34.py
import openpyxl
from openpyxl.chart import PieChart3D, Reference

wb = openpyxl.Workbook()                    # 開啟活頁簿
ws = wb.active                              # 獲得目前工作表
ws['A1']='五月份國外旅遊調查表'
rows=[['地點','大陸','東南亞','東北亞','美國','歐洲','紐澳'],
      ['人次',12000,18600,9600,7500,4860,5200]]
for row in rows:
    ws.append(row)

chart = PieChart3D()                          # 圓形圖
chart.title = '五月份國外旅遊調查表'

data = Reference(ws, min_col=1, min_row=3, max_col=7)      # 圖表資料
chart.add_data(data, titles_from_data=True) # 建立圖表
labels = Reference(ws, min_row=2, min_col = 2, max_col=7)   # 標籤名稱
chart.set_categories(labels)                # 設定標籤名稱
ws.add_chart(chart, 'A6')                   # 放置圖表位置
wb.save('ex7.xlsx')





